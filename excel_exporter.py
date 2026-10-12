"""
Excel Exporter Engine
Generates a multi-sheet Microsoft Excel workbook (.xlsx) using pandas and openpyxl.
Includes 6 standardized sheets: Android_Reviews, iOS_Reviews, Combined_Reviews,
App_Summary, Issue_Themes, and Collection_Log with tables, filters, frozen panes,
auto-fitted column widths, and wrapped text.
"""

from typing import List, Dict, Any, Optional
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd


# Color Palette
NAVY_HEADER = "1B365D"
SECTION_HEADER = "2C4C64"
WHITE_TEXT = "FFFFFF"
BORDER_GRAY = "D9D9D9"


def export_intelligence_workbook(
    android_reviews: List[Dict[str, Any]],
    ios_reviews: List[Dict[str, Any]],
    combined_reviews: List[Dict[str, Any]],
    metadata_list: List[Dict[str, Any]],
    summary_metrics: List[Dict[str, Any]],
    monthly_trends: List[Dict[str, Any]],
    segment_comparisons: Dict[str, Any],
    issue_themes: List[Dict[str, Any]],
    collection_logs: List[Dict[str, Any]],
    output_filepath: str,
) -> str:
    """
    Build the complete 6-sheet formatted Excel workbook.
    """
    wb = openpyxl.Workbook()
    # Remove default empty sheet
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)

    # Header style helper
    header_font = Font(name="Calibri", size=11, bold=True, color=WHITE_TEXT)
    header_fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    section_fill = PatternFill(start_color=SECTION_HEADER, end_color=SECTION_HEADER, fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color=BORDER_GRAY),
        right=Side(style="thin", color=BORDER_GRAY),
        top=Side(style="thin", color=BORDER_GRAY),
        bottom=Side(style="thin", color=BORDER_GRAY),
    )

    # -------------------------------------------------------------
    # 1. Sheet 1: Android_Reviews
    # -------------------------------------------------------------
    ws_android = wb.create_sheet(title="Android_Reviews")
    android_cols = [
        "App Name", "Store ID", "Platform", "App Role", "Review ID",
        "Reviewer Name", "Rating", "Review Text", "Review Date",
        "App Version", "Developer Reply", "Developer Reply Date",
        "Helpful Votes", "Collection Timestamp", "Source URL"
    ]
    ws_android.append(android_cols)
    for r in android_reviews:
        ws_android.append([
            r.get("app_name"), r.get("store_id"), r.get("platform"), r.get("app_role"),
            r.get("review_id"), r.get("reviewer_name"), r.get("rating"),
            r.get("review_text"), r.get("review_date"), r.get("app_version"),
            r.get("developer_reply"), r.get("developer_reply_date"),
            r.get("helpful_votes"), r.get("collection_timestamp"), r.get("source_url")
        ])

    # -------------------------------------------------------------
    # 2. Sheet 2: iOS_Reviews
    # -------------------------------------------------------------
    ws_ios = wb.create_sheet(title="iOS_Reviews")
    ios_cols = [
        "App Name", "Store ID", "Platform", "App Role", "Review ID",
        "Reviewer Name", "Rating", "Review Title", "Review Text", "Review Date",
        "App Version", "Developer Reply", "Developer Reply Date",
        "Helpful Votes", "Collection Timestamp", "Source URL"
    ]
    ws_ios.append(ios_cols)
    for r in ios_reviews:
        ws_ios.append([
            r.get("app_name"), r.get("store_id"), r.get("platform"), r.get("app_role"),
            r.get("review_id"), r.get("reviewer_name"), r.get("rating"),
            r.get("review_title"), r.get("review_text"), r.get("review_date"),
            r.get("app_version"), r.get("developer_reply"), r.get("developer_reply_date"),
            r.get("helpful_votes"), r.get("collection_timestamp"), r.get("source_url")
        ])

    # -------------------------------------------------------------
    # 3. Sheet 3: Combined_Reviews
    # -------------------------------------------------------------
    ws_combined = wb.create_sheet(title="Combined_Reviews")
    combined_cols = [
        "App Name", "Store ID", "Platform", "App Role", "Review ID",
        "Reviewer Name", "Rating", "Review Title", "Review Text", "Review Date",
        "App Version", "Developer Reply", "Developer Reply Date",
        "Helpful Votes", "Collection Timestamp", "Source URL"
